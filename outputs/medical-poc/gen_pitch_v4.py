#!/usr/bin/env python3
"""RingClaw Axis PoC Pitch v4 — entity memory + SOUL + personal/team bot"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK   = RGBColor(0x0D,0x1B,0x2A); CARD  = RGBColor(0x16,0x2A,0x40)
CARD2  = RGBColor(0x0F,0x22,0x35); BLUE  = RGBColor(0x00,0x8C,0xFF)
TEAL   = RGBColor(0x00,0xD4,0xBE); WHITE = RGBColor(0xFF,0xFF,0xFF)
GREY   = RGBColor(0xA8,0xBB,0xCC); DIM   = RGBColor(0x5A,0x7A,0x94)
GREEN  = RGBColor(0x00,0xC8,0x6E); RED   = RGBColor(0xFF,0x4D,0x4D)
YELLOW = RGBColor(0xFF,0xC1,0x07); PURP  = RGBColor(0x8B,0x5C,0xF6)
ORANGE = RGBColor(0xFF,0x7A,0x00)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
BL = prs.slide_layouts[6]

def S(): return prs.slides.add_slide(BL)
def BG(s,c=DARK):
    f=s.background.fill; f.solid(); f.fore_color.rgb=c
def R(s,x,y,w,h,c):
    sh=s.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background(); return sh
def T(s,txt,x,y,w,h,sz=14,bold=False,c=WHITE,align=PP_ALIGN.LEFT,italic=False,wrap=True):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tb.word_wrap=wrap; tf=tb.text_frame; tf.word_wrap=wrap
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=txt
    r.font.size=Pt(sz); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=c
def TM(s,lines,x,y,w,h,sz=12,c=WHITE,bold=False):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tb.word_wrap=True; tf=tb.text_frame; tf.word_wrap=True
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        r=p.add_run(); r.text=l; r.font.size=Pt(sz); r.font.bold=bold; r.font.color.rgb=c
def PILL(s,txt,x,y,w,h,bc=BLUE,fc=WHITE,sz=11):
    R(s,x,y,w,h,bc); T(s,txt,x+.06,y+.03,w-.12,h-.06,sz,True,fc,PP_ALIGN.CENTER)
def BAR(s,c=BLUE): R(s,0,0,.12,7.5,c)
def DIV(s,y,c=BLUE,x=.5,w=12.33):
    sh=s.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Pt(2))
    sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background()
def N(s,t): s.notes_slide.notes_text_frame.text=t

# ─── S1 Title ────────────────────────────────────────────────────────────────
s=S(); BG(s); BAR(s,BLUE)
PILL(s,"  Hackathon 2026  ",.5,.4,2.1,.36,TEAL,DARK,11)
T(s,"RingClaw",.5,.95,10,1.5,86,True)
T(s,"Your Team's AI — Built Inside RingEX",.5,2.55,11,.65,28,True,c=TEAL)
T(s,"Axis Integrated Mental Health  ·  Confirmed RC Customer  ·  Verified Data",.5,3.2,11,.42,15,c=GREY)
DIV(s,3.95)
T(s,'Charlotte (AIR) answers the calls  ·  RingClaw makes the clinic finish the job',.5,4.1,10,.38,14,italic=True,c=DIM)
for i,(l,c) in enumerate([("Personal Bot",BLUE),("Team Bot",TEAL),("Bot→Bot",GREEN)]):
    xi=9.9+i*1.1; R(s,xi,2.6,1,.9,c)
    T(s,l,xi-.05,3.58,1.12,.38,9,c=GREY,align=PP_ALIGN.CENTER)
N(s,"Charlotte（AIR）接了电话。RingClaw 让诊所把这件事做完。")

# ─── S2 Customer Profile ─────────────────────────────────────────────────────
s=S(); BG(s); BAR(s,TEAL)
PILL(s,"  真实 RC 付费客户（已验证）  ",.5,.35,3.2,.36,TEAL,DARK,11)
T(s,"Axis Integrated Mental Health",.5,.82,12,.7,34,True)

R(s,.5,1.68,5.85,5.42,CARD)
T(s,"机构概况（来自 axismh.com，深度验证）",.72,1.78,5.45,.38,13,True,c=TEAL)
facts=[("CEO","Christopher Perez（CRNA · 联合创始人）"),
       ("CMO/CGO","Liesl Perez（商业背景 · 联合创始人）"),
       ("CMO 2026","Dr. Kartiki Churi MD MBA FAPA"),
       ("诊所","4 个（Aurora · Westminster · Louisville · DTC）"),
       ("提供者","~18 人（9 PMHNP · 1 MD · 3 PA-C · 5 LPC/LCSW）"),
       ("患者量","1,000+（2023-12 下限）"),
       ("保险","99% 有保险（Medicaid · BCBS · Aetna · Cigna）")]
for i,(k,v) in enumerate(facts):
    yi=2.28+i*.68
    T(s,k+"：",.72,yi,1.7,.52,11,True,c=DIM)
    T(s,v,2.45,yi,3.75,.52,11)

R(s,6.55,1.68,6.3,2.6,CARD)
T(s,"已确认 RC 部署",6.77,1.78,5.9,.38,13,True,c=BLUE)
rc=[("AIR","Charlotte（AI 接待员名称）",GREEN),
    ("接通率","50% → 91%（来电 500→2,000+/周）",TEAL),
    ("自动化","30% 来电量自动处理",TEAL),
    ("EHR","athenahealth（2020 年上线）",GREY),
    ("续剂","电话 + 门户并存，行政仍是中间环节",YELLOW)]
for i,(k,v,c) in enumerate(rc):
    yi=2.28+i*.46; R(s,6.77,yi,.88,.34,c)
    T(s,k,6.79,yi+.04,.84,.26,10,True,c=DARK,align=PP_ALIGN.CENTER)
    T(s,v,7.72,yi+.02,4.95,.34,11)

R(s,6.55,4.42,6.3,2.68,CARD)
T(s,"Demo 真实提供者（axismh.com/providers/ 确认）",6.77,4.52,5.9,.38,13,True,c=YELLOW)
prov=[("Andrew Wenner","PMHNP · 开药（Scenario A + B 主角）",YELLOW),
      ("Michelle Godwin","LPC · 治疗师（同患者另一侧）",PURP),
      ("Alexis Gonzalez","行政协调员（续剂中间环节）",ORANGE),
      ("Christopher Perez","CEO（缺勤升级联系人）",GREY)]
for i,(name,role,c) in enumerate(prov):
    yi=4.96+i*.52
    T(s,name,6.77,yi,2.5,.42,12,True,c=c)
    T(s,role,9.35,yi,3.35,.42,11,c=GREY)
N(s,"所有提供者姓名均来自 axismh.com/providers/，深度研究 3-0 票验证。")

# ─── S3 Bot Architecture ─────────────────────────────────────────────────────
s=S(); BG(s); BAR(s,BLUE)
T(s,"在 RingEX 里创建 Bot",.5,.32,12.3,.65,34,True)
T(s,"每个员工都可以有自己的 Bot  ·  RingClaw 是那个 Bot  ·  SOUL.md 定义身份  ·  entity memory 是记忆",.5,.95,12.3,.42,16,c=GREY)

R(s,.5,1.52,5.85,5.58,CARD); R(s,.5,1.52,5.85,.5,BLUE)
T(s,"个人 Bot（只有本人可以触发）",.72,1.6,5.45,.35,14,True)
T(s,"员工为自己创建  ·  知道主人的患者、任务、偏好",.72,2.1,5.45,.35,12,c=GREY)
pbots=[("andrew-bot","Andrew Wenner, PMHNP",
        "· 知道 Andrew 的患者列表（entity memory）\n· 代表 Andrew 在临床频道发协调消息\n· 监听频道回复，路由临床反馈回 Andrew DM",YELLOW),
       ("alexis-bot","Alexis Gonzalez（行政）",
        "· 知道今日续剂任务队列（entity memory）\n· 缺勤时自动整理并向 nursecoord-bot 发交接信号\n· 向 Alexis DM 汇报交接结果",ORANGE)]
for i,(name,owner,desc,c) in enumerate(pbots):
    yi=2.6+i*2.12; R(s,.72,yi,5.45,1.88,CARD2); R(s,.72,yi,5.45,.38,c)
    T(s,name,.85,yi+.06,2,.26,12,True,c=DARK); T(s,owner,2.92,yi+.06,3.1,.26,11,c=DARK)
    T(s,desc,.85,yi+.48,5.1,1.3,11,wrap=True)

R(s,6.95,1.52,5.9,5.58,CARD); R(s,6.95,1.52,5.9,.5,TEAL)
T(s,"团队 Bot（频道里所有人都可以触发）",7.17,1.6,5.5,.35,14,True)
T(s,"团队共同创建  ·  处理团队级别的工作流",7.17,2.1,5.5,.35,12,c=GREY)
tbots=[("clinical-bot","#clinical-coordination",
        "· 接收续剂触发（人工 @bot 或 inbound SMS）\n· 读 entity memory 生成 Adaptive Card\n· 链式执行：SMS→患者 + TASK→Alexis",TEAL),
       ("nursecoord-bot","#admin",
        "· 接收 TASK_HANDOFF_REQUEST 信号\n· SMS 联系备用行政（Jennifer/Karen）\n· 确认后发 HANDOFF_COMPLETE 信号\n· 无回应时升级给 CEO Christopher",GREEN)]
for i,(name,ch,desc,c) in enumerate(tbots):
    yi=2.6+i*2.12; R(s,7.17,yi,5.45,1.88,CARD2); R(s,7.17,yi,5.45,.38,c)
    T(s,name,7.3,yi+.06,2,.26,12,True,c=DARK); T(s,ch,9.38,yi+.06,3.1,.26,11,c=DARK)
    T(s,desc,7.3,yi+.48,5.1,1.3,11,wrap=True)

DIV(s,7.18,BLUE)
T(s,"SOUL.md = Bot 的身份规则  ·  entity memory = Bot 的记忆  ·  DOMAIN.md = Bot 的知识库",.5,7.26,12.3,.2,12,c=DIM,align=PP_ALIGN.CENTER)
N(s,"展示 RingEX Bot 创建界面。每个 Bot 是一个 RingClaw 实例，--dir 隔离。")

# ─── S4 Entity Memory ────────────────────────────────────────────────────────
s=S(); BG(s); BAR(s,PURP)
PILL(s,"  entity memory  ",.5,.3,1.8,.36,PURP,DARK,11)
T(s,"Bot 的记忆：患者档案 + 任务队列 + 协调状态",2.45,.25,10.8,.55,24,True)
T(s,"Markdown 文件  ·  Bot 读取并写入  ·  每次操作后自动更新",.5,.92,12.3,.38,14,c=GREY)

# AX-2847
R(s,.5,1.45,3.95,5.65,CARD)
T(s,"AX-2847.md",.72,1.55,3.55,.38,13,True,c=TEAL)
T(s,"clinical-bot 读取  ·  Scenario A",.72,1.93,3.55,.32,10,c=DIM)
TM(s,["患者：Maria Lopez","手机：+17205550102","","用药：Sertraline 100mg","开药：Andrew Wenner PMHNP","上次就诊：05/28","上次续剂：04/15（51天前）","","保险：Blue Cross PPO ✅","药房：Walgreens Aurora","","─── 续剂记录 ───","2026-06-05 09:05","Andrew 批准 ✅","患者 SMS 已发 ✅","Alexis Task 创建 ✅"],.72,2.35,3.55,4.6,10)

# AX-1956
R(s,4.65,1.45,3.95,5.65,CARD)
T(s,"AX-1956.md",4.87,1.55,3.55,.38,13,True,c=YELLOW)
T(s,"andrew-bot 读取  ·  Scenario B",4.87,1.93,3.55,.32,10,c=DIM)
TM(s,["患者 ID：AX-1956","（隐去姓名，HIPAA 示范）","","主诊精神科：","  Andrew Wenner PMHNP","主诊治疗师：","  Michelle Godwin LPC","","用药：Sertraline 50mg","　→ 100mg（06/05 调整）","下次疗程：2026-06-12","","─── 协调记录 ───","14:31 变更通知发出","14:35 Michelle 回复","14:38 Andrew 知悉 ✅"],4.87,2.35,3.55,4.6,10)

# daily-queue
R(s,8.8,1.45,4.05,5.65,CARD)
T(s,"daily-queue-20260605.md",9.02,1.55,3.65,.38,12,True,c=ORANGE)
T(s,"alexis-bot 读取  ·  Scenario C",9.02,1.93,3.65,.32,10,c=DIM)
TM(s,["日期：2026-06-05","行政：Alexis Gonzalez","今日预约：24 名患者","","续剂 Task 队列（12 个）：","","✅ RX-047 Maria Lopez","   Sertraline 100mg","   Walgreens due 11:00","","⏳ RX-048 David M.","   Lexapro 20mg","   CVS due 12:00","","⏳ RX-052 Carlos M. ⚠️","   Adderall XR CONTROLLED","   due 13:00","（共 12 个…）"],9.02,2.35,3.65,4.6,10)

DIV(s,7.18,PURP)
T(s,"Bot 不需要 EHR 集成  ·  Demo 前预置文件  ·  Bot 运行中自动追加更新记录",.5,7.26,12.3,.2,12,c=DIM,align=PP_ALIGN.CENTER)
N(s,"Entity memory 是 Markdown 文件，存在 Bot 的 --dir 下。Bot 读它来生成 Card，写它来追踪状态。")

# ─── S5 Scenario A: Refill Routing ───────────────────────────────────────────
s=S(); BG(s); BAR(s,TEAL)
PILL(s,"  Scenario A · 团队 Bot  ",.5,.3,2.2,.36,TEAL,DARK,11)
T(s,"clinical-bot：续剂路由  （15-20 个/天 → 节省 10h 行政时间）",2.85,.25,10.4,.55,22,True)

R(s,.5,.98,12.3,.58,CARD2)
T(s,"Charlotte（AIR）接听 Maria 的续剂来电 → 发 SMS 确认 → Alexis 在 RingEX 看到记录",.7,1.04,11.9,.38,12,c=GREY)
T(s,'"@clinical-bot  refill AX-2847 Sertraline 100mg Andrew Wenner"',.7,1.38,11.9,.3,12,italic=True)

R(s,.5,1.85,6,3.5,CARD)
T(s,"clinical-bot 读 AX-2847.md → 生成 Card",.72,1.95,5.6,.35,12,True,c=TEAL)
T(s,"续剂请求 · RX-20260605-047",.72,2.35,5.6,.4,15,True)
rows=[("患者","Maria Lopez（AX-2847）"),("药物","Sertraline 100mg"),
      ("上次就诊","05/28（Andrew Wenner）"),("上次续剂","04/15（51 天前）✅ 合理"),("保险","Blue Cross PPO ✅")]
for i,(k,v) in enumerate(rows):
    yi=2.85+i*.36; T(s,k+"：",.72,yi,1.6,.3,11,c=DIM); T(s,v,2.35,yi,3.9,.3,12)
for j,(btn,c) in enumerate([("✅ 批准续剂",GREEN),("📅 需要问诊",YELLOW),("❌ 拒绝",RED)]):
    bx=.72+j*1.85; R(s,bx,4.95,1.68,.32,c)
    T(s,btn,bx+.04,4.97,1.6,.26,11,True,c=DARK,align=PP_ALIGN.CENTER)

T(s,"Andrew 点「✅ 批准续剂」",6.72,1.85,5.8,.35,13,c=GREY)
T(s,"↓  clinical-bot 链式执行",6.72,2.25,5.8,.32,12,c=DIM)
acts=[("ACTION:SMS","→ Maria +17205550102（即时）",TEAL),
      ("ACTION:TASK","→ Alexis：RX-047 发送至 Walgreens due 11:00",YELLOW),
      ("entity memory","→ AX-2847.md 追加续剂记录",GREY)]
for i,(a,r,c) in enumerate(acts):
    yi=2.72+i*.88; R(s,6.72,yi,1.62,.35,c)
    T(s,a,6.74,yi+.04,1.58,.27,10,True,c=DARK,align=PP_ALIGN.CENTER)
    T(s,r,8.4,yi+.04,4.15,.35,12)

R(s,6.72,5.4,2.68,.88,CARD2); T(s,"之前\n45–90 分钟",6.9,5.5,2.3,.68,16,True,c=RED)
R(s,9.55,5.4,2.95,.88,BLUE); T(s,"现在\n6 分钟",9.72,5.5,2.62,.68,18,True)
DIV(s,6.45)
T(s,"团队 Bot 的价值：任何人在频道触发 → 档案自动读取 → Card 发给正确提供者 → 链式执行",.5,6.55,12.3,.38,13,c=TEAL,align=PP_ALIGN.CENTER)
N(s,"重点：clinical-bot 从 entity memory 读档案，不需要查 EHR。提供者点一下，三件事同时发生。")

# ─── S6 Scenario B: Personal Bot Cross-provider ───────────────────────────────
s=S(); BG(s); BAR(s,YELLOW)
PILL(s,"  Scenario B · 个人 Bot  ",.5,.3,2.2,.36,YELLOW,DARK,11)
T(s,"andrew-bot：医生主动发起跨提供者协调（精神科独有问题）",2.85,.25,10.4,.55,22,True)

R(s,.5,.98,12.3,.5,CARD2)
T(s,"同一患者 AX-1956 同时有：Andrew Wenner PMHNP（开药）+ Michelle Godwin LPC（治疗）— 均为 Axis 真实在职提供者",.7,1.06,11.9,.32,12,c=GREY)

R(s,.5,1.65,5.8,1.45,CARD)
T(s,"Andrew 私信 andrew-bot（仅他可触发）",.72,1.75,5.4,.35,12,True,c=YELLOW)
T(s,'"AX-1956 今天调了 Sertraline\n 50mg→100mg，原因：抑郁改善不足。\n 通知 Michelle，让她关注焦虑波动。"',.72,2.18,5.4,.85,11,italic=True)

T(s,"→",6.42,2.18,.6,.55,22,c=DIM,align=PP_ALIGN.CENTER)

R(s,7.12,1.65,5.72,1.45,CARD)
T(s,"andrew-bot → #clinical-coordination  @Michelle Godwin",7.34,1.75,5.3,.35,12,True,c=YELLOW)
TM(s,["🔄 跨提供者用药更新 · AX-1956（Andrew，06/05）","Sertraline 50mg → 100mg","原因：抑郁症状改善不足","临床提示：SSRI 增量焦虑波动，下次疗程（06/12）重点评估","@Michelle Godwin 请回复确认收到"],7.34,2.18,5.3,.85,10)

R(s,.5,3.25,12.3,.5,CARD2)
T(s,"Michelle 自然语言回复（无需 @任何 Bot）：",.7,3.31,4.8,.32,12,c=DIM)
T(s,'"收到。她上次疗程（06/03）提到轻微心悸，Andrew 需要知道吗？"',5.52,3.31,7.1,.32,12,italic=True)

R(s,.5,3.9,5.8,1.55,CARD)
T(s,"andrew-bot 监听 #clinical-coordination",.72,4.0,5.4,.35,12,True,c=YELLOW)
TM(s,["匹配：发送者 = Michelle（watchlist 中）","内容：新临床症状（心悸）","判断：对开药医生相关 → 路由给 Andrew","→ coordination-watchlist.md 更新状态"],.72,4.42,5.4,.95,11)

T(s,"A→A\n→DM",6.42,4.1,.8,.9,14,True,c=GREEN,align=PP_ALIGN.CENTER)

R(s,7.12,3.9,5.72,1.55,CARD)
T(s,"andrew-bot → Andrew DM（私信）",7.34,4.0,5.3,.35,12,True,c=GREEN)
TM(s,["💬 Michelle 的临床反馈 · AX-1956","Michelle Godwin（14:35）：","患者 06/03 疗程提到轻微心悸","建议：随访时评估是否与","Sertraline 增量相关"],7.34,4.42,5.3,.95,11)

DIV(s,5.62,YELLOW)
for i,(w,d,c) in enumerate([
    ("EHR","记录系统\n不会主动路由信息",RED),
    ("AIR Charlotte","患者前台\n提供者间不经患者",DIM),
    ("andrew-bot","Andrew 的个人助理\nAndrew↔Michelle 信息闭环",GREEN)]):
    xi=.5+i*4.28; R(s,xi,5.72,4,1.38,CARD); R(s,xi,5.72,4,.42,c)
    T(s,w,xi+.12,5.79,3.76,.28,11,True,c=DARK); T(s,d,xi+.12,6.28,3.76,.75,11,c=WHITE,wrap=True)
N(s,"个人 Bot 是 Andrew 在 RingEX 里创建的。entity memory 里有双侧提供者信息，bot 知道 Michelle 是同一患者的治疗师。")

# ─── S7 Scenario C: Bot-to-Bot Handoff ───────────────────────────────────────
s=S(); BG(s); BAR(s,ORANGE)
PILL(s,"  Scenario C · Bot→Bot  ",.5,.3,2.2,.36,ORANGE,DARK,11)
T(s,"alexis-bot → nursecoord-bot：一句话缺勤，Bot 自动完成任务交接",2.85,.25,10.4,.55,22,True)

R(s,.5,.98,12.3,.5,CARD2)
T(s,"Alexis 是 Axis 行政核心——续剂 Task、患者跟进全经她。缺勤时 12 个续剂 Task 面临积压。",.7,1.06,11.9,.32,12,c=GREY)

steps=[
    ("Alexis\n→ alexis-bot","DM：\n'今天缺勤，帮我交接'",ORANGE,.5,1.62),
    ("alexis-bot\n处理","读 daily-queue-20260605.md\n续剂 12 个，最近 due 12:00\n发缺勤公告到 #admin",ORANGE,3.7,1.62),
    ("alexis-bot\n→ nursecoord-bot","#admin 发信号：\nTASK_HANDOFF_REQUEST\n续剂 12 个，due 12:00",GREEN,6.9,1.62),
    ("nursecoord-bot\n→ Jennifer SMS","ACTION:SMS +17205550301\n'12 refill tasks, due noon.\nCan you cover? Reply YES'",TEAL,10.1,1.62),
]
for (who,what,c,xi,yi) in steps:
    R(s,xi,yi,2.88,3.55,CARD); R(s,xi,yi,2.88,.55,c)
    T(s,who,xi+.12,yi+.1,2.65,.38,12,True,c=DARK)
    T(s,what,xi+.12,yi+.65,2.65,2.75,11,wrap=True)
    if xi<10: T(s,"→",xi+2.88,yi+1.5,.82,.5,18,c=DIM,align=PP_ALIGN.CENTER)

R(s,.5,5.35,12.3,.5,CARD2)
T(s,"Jennifer 回复 YES（6 分钟）→ nursecoord-bot 确认覆盖 → 发 HANDOFF_COMPLETE → alexis-bot 通知 Alexis DM ✅",.7,5.42,11.9,.32,12,c=GREEN)

R(s,.5,6.0,5.85,.78,CARD)
T(s,"升级路径（15 分钟无回应）",.72,6.08,5.45,.32,11,True,c=YELLOW)
T(s,"nursecoord-bot → ACTION:CARD → Christopher Perez（CEO）· 包含：试了谁、Task 数、预约量",.72,6.42,5.45,.3,10,c=GREY)
R(s,6.55,6.0,6.3,.78,CARD)
T(s,"Bot 间信号协议（#admin 频道消息总线）",6.77,6.08,5.9,.32,11,True,c=TEAL)
T(s,"TASK_HANDOFF_REQUEST → HANDOFF_COMPLETE → ESCALATED_TO_CEO",6.77,6.42,5.9,.3,10,c=GREY)
N(s,"Bot→Bot 通过 #admin 频道消息传递。nursecoord-bot 监听信号，alexis-bot 监听确认信号。无外部中间件。")

# ─── S8 Differentiation ───────────────────────────────────────────────────────
s=S(); BG(s); BAR(s,PURP)
T(s,"为什么是 RingEX + RingClaw",.5,.32,12.3,.65,32,True)

hdrs=["","MuleRun","RingCentral + RingClaw"]
cxs=[.5,4.5,8.5]; cws=[3.9,3.9,4.75]; ccs=[WHITE,GREY,TEAL]
for j,(h,x,w,c) in enumerate(zip(hdrs,cxs,cws,ccs)):
    if h: R(s,x,1.18,w,.45,CARD); T(s,h,x+.1,1.26,w-.2,.3,13,True,c=c,align=PP_ALIGN.CENTER)
rows2=[("Bot 创建方式","平台统一配置","员工在 RingEX 自助创建个人/团队 Bot"),
       ("个人 Bot","无","每员工有自己的 AI：知道自己的患者和任务"),
       ("entity memory","无","患者档案 + 任务状态，Bot 读写，自动追踪"),
       ("Bot 间协作","线程对话","个人 Bot ↔ 团队 Bot，信号协议，自动传递"),
       ("患者 SMS","IM 内部消息","双向 SMS 到患者手机（AIR + ACTION:SMS）"),
       ("医疗合规","通用 NLP","SOUL.md 硬规则：不提供医疗建议，不绕过审批")]
for i,(f,m,o) in enumerate(rows2):
    yi=1.75+i*.82; bg_r=CARD if i%2==0 else CARD2
    for j,(cell,x,w) in enumerate(zip([f,m,o],cxs,cws)):
        R(s,x,yi,w,.72,bg_r)
        c=WHITE
        if j==1: c=RED if any(kw in cell for kw in ["无","换","通用"]) else GREY
        if j==2: c=GREEN
        T(s,cell,x+.1,yi+.15,w-.2,.42,12,c=c,bold=(j==2),wrap=True)

DIV(s,6.68,PURP)
T(s,"Charlotte (AIR) 接了患者的电话  ·  每个员工有自己的 Bot  ·  RingEX 里完成业务闭环",.5,6.78,12.3,.38,13,True,align=PP_ALIGN.CENTER)
N(s,"最大差异：个人 Bot + entity memory。Andrew 的 Bot 知道他的患者，Alexis 的 Bot 知道她的任务。")

# ─── S9 Closing ───────────────────────────────────────────────────────────────
s=S(); BG(s); BAR(s,BLUE)
T(s,"Axis Integrated Mental Health  ·  真实 RC 客户  ·  真实提供者  ·  真实痛点",.5,.35,12.3,.58,22,True)
R(s,.5,1.08,12.3,.45,CARD)
T(s,"✅ CEO: Christopher Perez  ·  ✅ Charlotte（AIR，50%→91%）  ·  ✅ EHR: athenahealth  ·  ✅ Andrew Wenner + Michelle Godwin 已验证在职",.7,1.16,11.9,.28,11,c=GREEN)

scenes=[
    ("Scenario A","团队 Bot：clinical-bot 续剂路由","Alexis @bot 一行 → Card 给 Andrew → 链式执行\n每日节省 10 小时行政时间",TEAL),
    ("Scenario B","个人 Bot：andrew-bot 跨提供者协调","Andrew DM bot → bot 代发协调消息 → Michelle 回复 → bot 路由临床反馈回 Andrew DM",YELLOW),
    ("Scenario C","Bot→Bot：alexis-bot → nursecoord-bot","Alexis 一句话 → 12 个 Task 自动转移 → 6 分钟找到替班 → Alexis DM 收到汇报",ORANGE)]
for i,(label,title,desc,c) in enumerate(scenes):
    yi=1.72+i*1.58
    R(s,.5,yi,.88,1.32,c); T(s,label,.52,yi+.42,.84,.48,10,True,c=DARK,align=PP_ALIGN.CENTER)
    R(s,1.52,yi,11.3,1.32,CARD)
    T(s,title,1.72,yi+.12,5.5,.42,15,True,c=c)
    T(s,desc,1.72,yi+.58,10.9,.65,12,c=GREY,wrap=True)

DIV(s,6.52,BLUE)
T(s,"Charlotte 接了患者的电话  ·  RingClaw 让诊所里的每个人和他们的 Bot 把这件事做完",.5,6.62,12.3,.42,14,True,align=PP_ALIGN.CENTER)
N(s,"停顿：Charlotte 接了患者的电话。RingClaw 让诊所把这件事做完。")

# ─── Save ─────────────────────────────────────────────────────────────────────
out="/Users/summer.gan/Work/code/Aplication/ringclaw/outputs/medical-poc/RingClaw-Axis-PoC-v4.pptx"
prs.save(out); print(f"Saved: {out}  Slides: {len(prs.slides)}")
